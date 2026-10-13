from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


TIMESTAMP_FORMATS = ("%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %H:%M:%S")


@dataclass(frozen=True)
class AttackInterval:
    scenario_id: int
    start: datetime
    end: datetime
    attack_point: str
    attack: str

    def to_dict(self) -> dict:
        result = asdict(self)
        result["start"] = self.start.isoformat()
        result["end"] = self.end.isoformat()
        return result


@dataclass(frozen=True)
class EventGroup:
    event_id: str
    start: datetime
    end: datetime
    scenario_ids: tuple[int, ...]
    source: str

    @property
    def duration_seconds(self) -> int:
        return int((self.end - self.start).total_seconds()) + 1

    def to_dict(self) -> dict:
        return {
            "event_id": self.event_id,
            "start": self.start.isoformat(),
            "end": self.end.isoformat(),
            "duration_seconds": self.duration_seconds,
            "scenario_ids": list(self.scenario_ids),
            "source": self.source,
        }


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_timestamp(value: object) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    raw = str(value).strip()
    for timestamp_format in TIMESTAMP_FORMATS:
        try:
            return datetime.strptime(raw, timestamp_format)
        except ValueError:
            continue
    raise ValueError(f"Unsupported timestamp: {value!r}")


def is_attack_label(value: object) -> bool:
    return "attack" in re.sub(r"\s+", "", str(value).lower())


def _find_header(worksheet) -> tuple[int, list[str]]:
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True),
        start=1,
    ):
        header = ["" if value is None else str(value).strip() for value in row]
        if "Timestamp" in header and "Normal/Attack" in header:
            return row_index, header
    raise ValueError(f"Could not find SWaT header in worksheet {worksheet.title!r}")


def iter_swat_rows(path: Path) -> tuple[list[str], Iterator[tuple]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header_row, header = _find_header(worksheet)

    def rows() -> Iterator[tuple]:
        try:
            yield from worksheet.iter_rows(min_row=header_row + 1, values_only=True)
        finally:
            workbook.close()

    return header, rows()


def _normalized_start(value: object) -> datetime:
    start = parse_timestamp(value)
    # The provider attack table records 2 January as 2015 for scenarios 37-41,
    # while the continuous telemetry and scenario order place them in 2016.
    if start.year == 2015 and start.month == 1:
        start = start.replace(year=2016)
    return start


def _combine_end(start: datetime, value: object) -> datetime:
    if isinstance(value, datetime):
        end = value.replace(tzinfo=None)
    elif isinstance(value, time):
        end = datetime.combine(start.date(), value)
    else:
        parsed = None
        raw = str(value).strip()
        for timestamp_format in ("%H:%M:%S", "%I:%M:%S %p"):
            try:
                parsed = datetime.strptime(raw, timestamp_format).time()
                break
            except ValueError:
                continue
        if parsed is None:
            raise ValueError(f"Unsupported interval end: {value!r}")
        end = datetime.combine(start.date(), parsed)
    return end + timedelta(days=1) if end < start else end


def load_attack_intervals(path: Path) -> tuple[list[AttackInterval], dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    indexes = {name: header.index(name) for name in ("Attack #", "Start Time", "End Time")}
    point_index = header.index("Attack Point")
    attack_index = header.index("Attack")
    intervals: list[AttackInterval] = []
    excluded: list[dict] = []
    corrected_years: list[int] = []
    for row in rows:
        raw_id = row[indexes["Attack #"]]
        if not isinstance(raw_id, (int, float)):
            continue
        scenario_id = int(raw_id)
        start_value = row[indexes["Start Time"]]
        end_value = row[indexes["End Time"]]
        attack_point = str(row[point_index] or "").strip()
        if end_value is None or "no physical impact" in attack_point.lower():
            excluded.append({"scenario_id": scenario_id, "reason": "no_completed_physical_interval"})
            continue
        original_start = parse_timestamp(start_value)
        start = _normalized_start(start_value)
        if start.year != original_start.year:
            corrected_years.append(scenario_id)
        intervals.append(
            AttackInterval(
                scenario_id=scenario_id,
                start=start,
                end=_combine_end(start, end_value),
                attack_point=attack_point,
                attack=str(row[attack_index] or "").strip(),
            )
        )
    workbook.close()
    intervals.sort(key=lambda interval: interval.start)
    metadata = {
        "usable_interval_count": len(intervals),
        "excluded_rows": excluded,
        "corrected_year_scenarios": corrected_years,
    }
    return intervals, metadata


def audit_swat_workbook(path: Path, intervals: list[AttackInterval] | None = None) -> dict:
    header, rows = iter_swat_rows(path)
    timestamp_index = header.index("Timestamp")
    label_index = header.index("Normal/Attack")
    labels: Counter[str] = Counter()
    normalized_labels: Counter[str] = Counter()
    seen: set[datetime] = set()
    missing_by_column: Counter[str] = Counter()
    previous: datetime | None = None
    minimum: datetime | None = None
    maximum: datetime | None = None
    duplicates = regressions = parse_failures = gaps = row_count = 0
    gap_events: list[dict] = []
    attack_label_runs: list[dict] = []
    active_label_run_start: datetime | None = None
    active_label_run_end: datetime | None = None
    interval_but_normal = labeled_attack_outside_interval = 0
    scenario_rows: Counter[int] = Counter()
    interval_index = 0
    ordered_intervals = intervals or []

    for row in rows:
        if not row or all(value is None for value in row):
            continue
        row_count += 1
        for column_index, column_name in enumerate(header):
            if column_index >= len(row) or row[column_index] is None:
                missing_by_column[column_name] += 1
        label = str(row[label_index]).strip()
        labels[label] += 1
        is_attack = is_attack_label(label)
        normalized_labels["Attack" if is_attack else "Normal"] += 1
        try:
            timestamp = parse_timestamp(row[timestamp_index])
        except ValueError:
            parse_failures += 1
            continue
        if timestamp in seen:
            duplicates += 1
        seen.add(timestamp)
        if previous is not None:
            if timestamp < previous:
                regressions += 1
            elif timestamp - previous > timedelta(seconds=1):
                gaps += 1
                gap_events.append(
                    {
                        "previous": previous.isoformat(),
                        "current": timestamp.isoformat(),
                        "seconds": int((timestamp - previous).total_seconds()),
                    }
                )
        previous = timestamp
        minimum = timestamp if minimum is None or timestamp < minimum else minimum
        maximum = timestamp if maximum is None or timestamp > maximum else maximum

        while interval_index < len(ordered_intervals) and timestamp > ordered_intervals[interval_index].end:
            interval_index += 1
        active = (
            ordered_intervals[interval_index]
            if interval_index < len(ordered_intervals)
            and ordered_intervals[interval_index].start <= timestamp
            else None
        )
        if is_attack and active_label_run_start is None:
            active_label_run_start = timestamp
            active_label_run_end = timestamp
        elif is_attack:
            active_label_run_end = timestamp
        elif not is_attack and active_label_run_start is not None:
            attack_label_runs.append(
                {
                    "start": active_label_run_start.isoformat(),
                    "end": active_label_run_end.isoformat(),
                }
            )
            active_label_run_start = None
            active_label_run_end = None
        if active is not None:
            scenario_rows[active.scenario_id] += 1
            if not is_attack:
                interval_but_normal += 1
        elif is_attack:
            labeled_attack_outside_interval += 1

    if active_label_run_start is not None and active_label_run_end is not None:
        attack_label_runs.append(
            {"start": active_label_run_start.isoformat(), "end": active_label_run_end.isoformat()}
        )

    return {
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "worksheet_rows_including_headers": row_count + 2,
        "rows": row_count,
        "columns": len(header),
        "column_names": header,
        "label_counts": dict(labels),
        "normalized_label_counts": dict(normalized_labels),
        "unique_timestamps": len(seen),
        "duplicate_timestamps": duplicates,
        "chronology_regressions": regressions,
        "timestamp_gaps": gaps,
        "timestamp_gap_events": gap_events,
        "timestamp_parse_failures": parse_failures,
        "timestamp_min": minimum.isoformat() if minimum else None,
        "timestamp_max": maximum.isoformat() if maximum else None,
        "missing_by_column": dict(missing_by_column),
        "attack_label_runs": attack_label_runs,
        "interval_label_consistency": {
            "interval_but_labeled_normal": interval_but_normal,
            "labeled_attack_outside_interval": labeled_attack_outside_interval,
            "rows_by_scenario": dict(sorted(scenario_rows.items())),
        }
        if intervals is not None
        else None,
    }


def classify_official_dataset(
    normal: dict, attack: dict, intervals: list[AttackInterval], headers_match: bool
) -> dict:
    blocking: list[str] = []
    warnings: list[str] = []
    for name, item in (("normal", normal), ("attack", attack)):
        for field in ("duplicate_timestamps", "chronology_regressions", "timestamp_parse_failures"):
            if item[field]:
                blocking.append(f"{name}: {field}={item[field]}")
        if item["missing_by_column"]:
            blocking.append(f"{name}: missing values detected")
        if item["timestamp_gaps"]:
            warnings.append(f"{name}: timestamp_gaps={item['timestamp_gaps']}")
    if not headers_match:
        blocking.append("normal and attack feature schemas differ")
    if set(normal["label_counts"]) != {"Normal"}:
        blocking.append("normal workbook contains unexpected labels")
    if not {"Normal", "Attack"}.issubset(attack["label_counts"]):
        blocking.append("attack workbook is not a continuous mixed-label sequence")
    if len(intervals) != 36:
        blocking.append(f"expected 36 completed physical attack intervals, found {len(intervals)}")
    consistency = attack["interval_label_consistency"]
    if consistency and (
        consistency["interval_but_labeled_normal"]
        or consistency["labeled_attack_outside_interval"]
    ):
        warnings.append("provider labels and attack-table intervals are not identical")
    return {
        "gate1_pass": not blocking,
        "publication_data_accepted": not blocking,
        "blocking_findings": blocking,
        "warnings": warnings,
    }


def reconcile_event_groups(
    intervals: list[AttackInterval], label_runs: list[dict]
) -> list[EventGroup]:
    groups: list[EventGroup] = []
    mapped_scenarios: set[int] = set()
    for index, run in enumerate(label_runs, start=1):
        start = datetime.fromisoformat(run["start"])
        end = datetime.fromisoformat(run["end"])
        overlapping = tuple(
            interval.scenario_id
            for interval in intervals
            if max(start, interval.start) <= min(end, interval.end)
        )
        mapped_scenarios.update(overlapping)
        groups.append(
            EventGroup(
                event_id=f"event_{index:02d}",
                start=start,
                end=end,
                scenario_ids=overlapping,
                source="official_intervals" if overlapping else "provider_label_only",
            )
        )
    expected = {interval.scenario_id for interval in intervals}
    missing = sorted(expected - mapped_scenarios)
    if missing:
        raise ValueError(f"Official scenarios are not represented by provider labels: {missing}")
    return groups


def build_event_blocks(
    events: list[EventGroup], dataset_start: datetime, dataset_end: datetime
) -> list[dict]:
    if not events or dataset_start > events[0].start or dataset_end < events[-1].end:
        raise ValueError("Dataset bounds must contain every event")
    blocks: list[dict] = []
    for index, event in enumerate(events):
        if index == 0:
            block_start = dataset_start
        else:
            previous = events[index - 1]
            midpoint = previous.end + (event.start - previous.end) / 2
            block_start = midpoint.replace(microsecond=0) + timedelta(seconds=1)
        if index == len(events) - 1:
            block_end = dataset_end
        else:
            following = events[index + 1]
            block_end = (event.end + (following.start - event.end) / 2).replace(microsecond=0)
        blocks.append(
            {
                **event.to_dict(),
                "block_start": block_start.isoformat(),
                "block_end": block_end.isoformat(),
                "block_seconds": int((block_end - block_start).total_seconds()) + 1,
            }
        )
    return blocks


def assign_balanced_folds(events: list[EventGroup], fold_count: int = 5) -> dict[str, int]:
    if fold_count < 2 or len(events) < fold_count:
        raise ValueError("fold_count must permit at least one event per fold")
    ordered = sorted(events, key=lambda event: (-event.duration_seconds, event.event_id))
    snake = list(range(fold_count)) + list(range(fold_count - 1, -1, -1))
    return {event.event_id: snake[index % len(snake)] for index, event in enumerate(ordered)}
